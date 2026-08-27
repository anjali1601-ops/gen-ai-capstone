"""Create .docx and .pdf sample complaints from the text files."""

from pathlib import Path

from docx import Document
from fpdf import FPDF

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"
SOURCES = Path(__file__).resolve().parent / "sources"


def write_docx(source: Path, target: Path) -> None:
    document = Document()
    document.add_heading("Customer Complaint", level=1)
    for line in source.read_text(encoding="utf-8").splitlines():
        document.add_paragraph(line)
    document.save(target)


def write_pdf(source: Path, target: Path) -> None:
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_left_margin(15)
    pdf.set_right_margin(15)
    pdf.set_font("Helvetica", size=11)
    usable_width = pdf.w - pdf.l_margin - pdf.r_margin
    for line in source.read_text(encoding="utf-8").splitlines():
        safe = line.encode("latin-1", "replace").decode("latin-1")
        if not safe.strip():
            pdf.ln(6)
            continue
        # fpdf2 defaults leave the cursor at the right edge; reset to the left margin.
        pdf.multi_cell(usable_width, 6, safe, new_x="LMARGIN", new_y="NEXT")
    pdf.output(str(target))


def main() -> None:
    DATA.mkdir(parents=True, exist_ok=True)
    write_docx(SOURCES / "complaint_003.txt", DATA / "complaint_003.docx")
    write_pdf(SOURCES / "complaint_004.txt", DATA / "complaint_004.pdf")
    write_pdf(SOURCES / "complaint_005.txt", DATA / "complaint_005.pdf")
    print("Wrote sample .docx and .pdf files in data/")


if __name__ == "__main__":
    main()
